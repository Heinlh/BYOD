"""Authored cross-topic documents and labeled natural-language retrieval checks."""

import json
from pathlib import Path

import pymupdf
from docx import Document
from pptx import Presentation

DESTINATION = Path(__file__).resolve().parents[1] / "tests" / "fixtures" / "queries"

ECONOMICS = [
    (
        "Opportunity cost",
        "Opportunity cost is the value of the best alternative given up when a "
        "choice is made. A student who spends an evening studying instead of working gives up "
        "the wages from that shift. The opportunity cost is those forgone wages, not every "
        "possible activity added together. Scarcity forces people to choose between competing "
        "uses of limited time and resources.",
    ),
    (
        "Price elasticity",
        "Price elasticity of demand measures how strongly the quantity demanded "
        "responds to a change in price. It equals the percentage change in quantity demanded "
        "divided by the percentage change in price. Demand is elastic when the absolute value "
        "exceeds one. When demand is elastic, raising the price reduces total revenue because "
        "the fall in sales volume outweighs the higher price per unit.",
    ),
    (
        "Inflation",
        "Inflation is a sustained increase in the general price level. It reduces the "
        "purchasing power of money: the same amount of currency buys fewer goods and services. "
        "The consumer price index tracks the cost of a representative basket over time. "
        "An increase in one product's price alone is not necessarily inflation across the economy. "
        "Real wages adjust nominal wages for changes in the general price level.",
    ),
]
STATISTICS = [
    (
        "Confidence intervals",
        "A 95 percent confidence interval comes from a procedure that would "
        "contain the fixed population parameter in 95 percent of repeated samples. The parameter "
        "does not change between repetitions; the interval endpoints do. A larger sample usually "
        "reduces standard error and narrows the interval, holding variability and the confidence "
        "level constant. Greater confidence requires a wider interval.",
    ),
    (
        "P values",
        "A p value is the probability, assuming the null hypothesis is true, of observing "
        "a result at least as extreme as the one obtained. It is not the probability that the null "
        "hypothesis is true. A small p value suggests that the observations are unusual under "
        "the null model. Statistical significance alone does not establish a large effect or "
        "practical importance.",
    ),
    (
        "Random assignment",
        "Random assignment places study participants into treatment groups "
        "using chance. It helps balance both measured and unmeasured confounding factors on "
        "average, supporting causal comparisons between treatments. Random sampling instead "
        "selects people from a population and supports generalization. Assignment and sampling "
        "solve different problems; a randomly assigned experiment need not use a random "
        "sample of the population.",
    ),
]
NETWORKS = [
    (
        "TCP reliability",
        "TCP provides reliable, ordered delivery of a byte stream. Sequence "
        "numbers identify the position of data in the stream. Acknowledgments tell the sender "
        "which data arrived, and missing acknowledgments trigger retransmission. The receiver "
        "reorders segments before delivering bytes to the application. TCP also controls flow "
        "so a fast sender does not overwhelm a slow receiver.",
        "",
    ),
    (
        "DNS lookup",
        "DNS translates a domain name into an IP address. A resolver checks its cache "
        "before asking other DNS servers. If a cached record is still valid, it can answer "
        "without a new remote lookup. Otherwise, recursive resolution follows the DNS hierarchy "
        "until an authoritative server supplies the record.",
        "The time to live, or TTL, controls how long a DNS answer may remain cached. A changed "
        "website address may appear stale until the old cached record's TTL expires. Lowering "
        "TTL before a planned migration reduces the duration of stale cached answers.",
    ),
    (
        "HTTPS encryption",
        "HTTPS carries HTTP over Transport Layer Security, abbreviated TLS. "
        "The TLS handshake authenticates the server using its certificate and establishes "
        "shared session keys. Encryption protects application data from passive eavesdroppers "
        "while integrity checks reveal tampering in transit. HTTPS does not guarantee that "
        "the website's claims or business practices are trustworthy.",
        "",
    ),
    (
        "UDP tradeoffs",
        "UDP sends independent datagrams without TCP's reliable ordered stream. "
        "It does not automatically retransmit lost packets or reorder arriving packets for "
        "the application. This reduces protocol overhead and can suit time-sensitive voice "
        "or video. Applications that need reliability over UDP must supply that behavior "
        "themselves. Lower latency is a tradeoff, not a promise that packets cannot be lost.",
        "",
    ),
]


def main() -> None:
    DESTINATION.mkdir(parents=True, exist_ok=True)
    with pymupdf.open() as pdf:
        for title, body in ECONOMICS:
            page = pdf.new_page()
            page.insert_text((50, 60), title, fontsize=22)
            page.insert_textbox(pymupdf.Rect(50, 100, 540, 700), body, fontsize=12)
        pdf.save(DESTINATION / "economics.pdf")
    doc = Document()
    for title, body in STATISTICS:
        doc.add_heading(title, 1)
        doc.add_paragraph(body)
    doc.save(str(DESTINATION / "statistics.docx"))
    deck = Presentation()
    deck.core_properties.title = "Computer Networks"
    for title, body, notes in NETWORKS:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        if notes:
            frame = slide.notes_slide.notes_text_frame
            if frame is not None:
                frame.text = notes
    deck.save(str(DESTINATION / "networks.pptx"))
    cases = [
        (
            "What am I giving up if I study tonight instead of taking a paid shift?",
            "economics.pdf",
            "p. 1",
            "forgone wages",
        ),
        (
            "Why could charging more make a business earn less revenue?",
            "economics.pdf",
            "p. 2",
            "elastic",
        ),
        (
            "What happens to the buying power of cash when prices rise generally?",
            "economics.pdf",
            "p. 3",
            "purchasing power",
        ),
        (
            "What does a 95 percent confidence interval mean over repeated samples?",
            "statistics.docx",
            "Confidence intervals",
            "repeated samples",
        ),
        (
            "If I collect a bigger sample, what happens to the uncertainty interval?",
            "statistics.docx",
            "Confidence intervals",
            "narrows",
        ),
        (
            "Is a p value the chance that my null hypothesis is true?",
            "statistics.docx",
            "P values",
            "not the probability",
        ),
        (
            "How does assigning people to treatments by chance help with confounding?",
            "statistics.docx",
            "Random assignment",
            "confounding",
        ),
        (
            "How can TCP recover data that never reaches the receiver?",
            "networks.pptx",
            "slide 1",
            "retransmission",
        ),
        (
            "Why might a website still resolve to its old address after a DNS change?",
            "networks.pptx",
            "slide 2",
            "TTL",
        ),
        (
            "How does HTTPS stop someone on the network from reading application data?",
            "networks.pptx",
            "slide 3",
            "Encryption",
        ),
        (
            "Why would a video call use UDP even though some packets can disappear?",
            "networks.pptx",
            "slide 4",
            "voice or video",
        ),
    ]
    manifest = {
        "supported": [
            {"query": q, "filename": f, "locator": loc, "evidence": evidence}
            for q, f, loc, evidence in cases
        ],
        "unsupported": [
            "Who wrote Pride and Prejudice?",
            "How do I bake a chocolate cake?",
            "Who won the football world cup in 1998?",
            "What is tomorrow's weather?",
        ],
    }
    (DESTINATION / "cases.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
