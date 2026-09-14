"""Repeatable scale workspaces and the latency matrix for `byod debug bench`.

Documents are real PDF, DOCX, and PPTX files assembled from docstring prose installed in
BYOD's own locked environment (offline and deterministic for that environment). They are
ingested through the normal parser, chunker, and ONNX pipeline, and the workspace is
snapshotted at 1k, 5k, and 25k chunks. The labeled query fixtures are included so every
tier has known answer documents for scoped runs.

Local embedding measured 1.2-1.6 chunks/s on the development laptop, so a fully embedded
25k corpus takes hours. Beyond --real-chunks (default 5000) the build copies existing
indexed documents: each copy gets its own document row, chunks, and the true vectors of
its text. Pass --real-chunks 25000 for a fully embedded corpus.

    uv run python scripts/bench_scale.py build
    uv run python scripts/bench_scale.py run --runs 20
    uv run python scripts/bench_scale.py run --runs 1 --cases 7 --llm qwen3.6:latest

Output contains durations and counts only. Data lives under --root, outside the repository.
"""

import argparse
import ast
import json
import os
import random
import shutil
import sqlite3
import statistics
import subprocess
import sys
import sysconfig
import time
from contextlib import closing
from pathlib import Path

import onnxruntime as ort
import pymupdf
from docx import Document
from platformdirs import user_cache_path
from pptx import Presentation

from byod.config import Config, Selection
from byod.db.migrations import connect, initialize
from byod.db.queries import create_workspace
from byod.index.embed import Embedder
from byod.ingest.jobs import JobQueue
from byod.model_files import MODEL_REVISION, artifact

REPO = Path(__file__).resolve().parents[1]
QUERIES = REPO / "tests" / "fixtures" / "queries"
DEV = Config(REPO / ".byod-dev")
TIERS = (1000, 5000, 25000)
WORKSPACE = "bench"


def paragraphs() -> list[str]:
    paths = sysconfig.get_paths()
    found: list[str] = []
    for root in dict.fromkeys((paths["stdlib"], paths["purelib"])):
        for file in sorted(Path(root).rglob("*.py")):
            skip = {"test", "tests", "idlelib", "site-packages", "_vendor"}
            if skip & set(file.relative_to(root).parts):
                continue
            try:
                tree = ast.parse(file.read_bytes())
            except (SyntaxError, ValueError):
                continue
            for node in ast.walk(tree):
                if isinstance(node, ast.Module | ast.ClassDef | ast.FunctionDef):
                    for part in (ast.get_docstring(node) or "").split("\n\n"):
                        text = " ".join(part.split())
                        if len(text) >= 200 and ">>>" not in text and text.isascii():
                            found.append(text)
    return found


def passage(rng: random.Random, pool: list[str], anchor: int, size: int) -> str:
    # Neighbouring docstrings keep a passage on one topic while varying its composition.
    text = ""
    while len(text) < size:
        text += " " + pool[(anchor + rng.randrange(40)) % len(pool)]
    return text.strip()[: size + 400].rsplit(" ", 1)[0]


def write_document(index: int, pool: list[str], folder: Path) -> Path:
    rng = random.Random(index)  # noqa: S311 - deterministic benchmark content
    anchor = rng.randrange(len(pool))
    kind = ("pdf", "pdf", "docx", "pptx")[index % 4]
    path = folder / f"synthetic-{index:04d}.{kind}"

    def title() -> str:
        return " ".join(pool[(anchor + rng.randrange(40)) % len(pool)].split()[:5])

    if kind == "pdf":
        with pymupdf.open() as pdf:
            for number in range(rng.randint(40, 200)):
                page = pdf.new_page()
                top = 60
                if number % 6 == 0:
                    page.insert_text((50, 60), title(), fontsize=18)
                    top = 100
                body = passage(rng, pool, anchor, 1800)
                if page.insert_textbox(pymupdf.Rect(50, top, 545, 790), body, fontsize=10) < 0:
                    raise RuntimeError("Generated page text did not fit")
                anchor += rng.randrange(3)
            pdf.save(path)
    elif kind == "docx":
        document = Document()
        for _ in range(rng.randint(10, 40)):
            document.add_heading(title(), rng.randint(1, 3))
            for _ in range(rng.randint(1, 4)):
                document.add_paragraph(passage(rng, pool, anchor, 700))
            anchor += rng.randrange(3)
        document.save(str(path))
    else:
        deck = Presentation()
        for number in range(rng.randint(20, 60)):
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            if slide.shapes.title is not None:
                slide.shapes.title.text = title()
            slide.placeholders[1].text = passage(rng, pool, anchor, 600)
            if number % 2 == 0:
                slide.notes_slide.notes_text_frame.text = passage(rng, pool, anchor, 400)
            anchor += rng.randrange(3)
        deck.save(str(path))
    return path


def tier_config(root: Path, target: int) -> Config:
    config = Config(root / f"{target // 1000}k")
    folder = config.data_dir / "models" / "jina-embeddings-v2-base-en" / MODEL_REVISION
    folder.mkdir(parents=True, exist_ok=True)
    for source in (artifact(DEV), artifact(DEV, weights=True)):
        if not (folder / source.name).exists():
            try:
                os.link(source, folder / source.name)
            except OSError:
                shutil.copyfile(source, folder / source.name)
    return config


def chunk_total(config: Config) -> int:
    with connect(config) as db:
        return int(db.execute("SELECT count(*) FROM chunks").fetchone()[0])


def ingest(jobs: JobQueue, workspace: int, paths: list[str]) -> None:
    identifier = jobs.submit(workspace, paths)
    jobs.wait()
    status = jobs.snapshot(identifier)
    if not status or status["state"] != "done":
        raise SystemExit(f"Ingest failed: {status and status['errors']}")


def replicate(config: Config, target: int) -> None:
    with connect(config) as db:
        originals = [
            int(row[0])
            for row in db.execute(
                "SELECT id FROM documents WHERE filename LIKE 'synthetic-%' "
                "AND status='indexed' ORDER BY id"
            )
        ]
        total = int(db.execute("SELECT count(*) FROM chunks").fetchone()[0])
        copy = 1 + int(
            db.execute("SELECT count(*) FROM documents WHERE filename LIKE 'copy%'").fetchone()[0]
        )
        while total < target:
            source = originals[(copy - 1) % len(originals)]
            document = db.execute(
                "INSERT INTO documents(workspace_id,filename,path,doc_type,content_hash,unit_count,"
                "chunk_count,status,embed_model,embed_version,ingested_at) "
                "SELECT workspace_id,'copy'||?||'-'||filename,path,doc_type,content_hash||'-'||?,"
                "unit_count,chunk_count,status,embed_model,embed_version,ingested_at "
                "FROM documents WHERE id=?",
                (copy, copy, source),
            ).lastrowid
            for row in db.execute(
                "SELECT id,ordinal,text,locator,block_type,token_count FROM chunks "
                "WHERE document_id=? ORDER BY ordinal",
                (source,),
            ).fetchall():
                chunk = db.execute(
                    "INSERT INTO chunks(document_id,ordinal,text,locator,block_type,token_count) "
                    "VALUES (?,?,?,?,?,?)",
                    (document, *tuple(row)[1:]),
                ).lastrowid
                db.execute(
                    "INSERT INTO chunk_vectors(chunk_id,embedding) "
                    "SELECT ?,embedding FROM chunk_vectors WHERE chunk_id=?",
                    (chunk, row[0]),
                )
                total += 1
            copy += 1


def build(root: Path, real_chunks: int) -> None:
    pool = paragraphs()
    sources = root / "sources"
    sources.mkdir(parents=True, exist_ok=True)
    embedder = Embedder(DEV)
    # Build-time only: embed with every core. The application keeps its own session.
    options = ort.SessionOptions()
    options.intra_op_num_threads = os.cpu_count() or 4
    options.log_severity_level = 4
    embedder._session = ort.InferenceSession(
        str(artifact(DEV, weights=True)), sess_options=options, providers=["CPUExecutionProvider"]
    )
    previous: Config | None = None
    for target in TIERS:
        config = tier_config(root, target)
        if previous and not config.database.exists():
            with (
                closing(sqlite3.connect(previous.database)) as source,
                closing(sqlite3.connect(config.database)) as copy,
            ):
                source.backup(copy)
        initialize(config)
        with connect(config) as db:
            row = db.execute("SELECT id FROM workspaces WHERE name=?", (WORKSPACE,)).fetchone()
            workspace = int(row[0]) if row else create_workspace(db, WORKSPACE)
        jobs = JobQueue(config, embedder.index_document)
        try:
            ingest(jobs, workspace, [str(QUERIES)])
            while chunk_total(config) < target:
                if chunk_total(config) >= real_chunks:
                    replicate(config, target)
                    break
                with connect(config) as db:
                    index = db.execute(
                        "SELECT count(*) FROM documents WHERE filename LIKE 'synthetic-%'"
                    ).fetchone()[0]
                started = time.perf_counter()
                path = write_document(int(index), pool, sources)
                ingest(jobs, workspace, [str(path)])
                seconds = round(time.perf_counter() - started, 1)
                print(
                    json.dumps({"tier": target, "chunks": chunk_total(config), "s": seconds}),
                    flush=True,
                )
        finally:
            jobs.close()
        previous = config


def run(root: Path, runs: int, cases: list[int], model: str | None) -> None:
    supported = json.loads((QUERIES / "cases.json").read_text(encoding="utf-8"))["supported"]
    reports: list[dict[str, object]] = []
    for target in TIERS:
        config = tier_config(root, target)
        if model:
            with config.edit_preferences() as preferences:
                preferences.active = Selection(provider="ollama", model=model)
        with connect(config) as db:
            ids = {
                r["filename"]: int(r["id"]) for r in db.execute("SELECT id,filename FROM documents")
            }
        total = chunk_total(config)
        for case in cases:
            filename = supported[case]["filename"]
            for extra in ([], ["--doc", str(ids[filename])]):
                command = [
                    sys.executable, "-m", "byod.cli", "--data-dir", str(config.data_dir),
                    "debug", "bench", WORKSPACE, supported[case]["query"],
                    "--runs", str(runs), *extra, *([] if model else ["--no-llm"]),
                ]  # fmt: skip
                completed = subprocess.run(command, capture_output=True, check=True, text=True)  # noqa: S603
                report = json.loads(completed.stdout)
                report.update(workspace_chunks=total, case=case)
                print(json.dumps(report), flush=True)
                reports.append(report)
    table(reports)


def table(reports: list[dict[str, object]]) -> None:
    """Median of per-query medians; p95 column is the worst per-query p95."""
    phases = ("candidates", "embed", "search", "assembly", "retrieval", "ttft")
    print("\n| chunks | scope | candidates | ctx chunks | ctx tokens | " + " | ".join(
        f"{p} ms med / p95" for p in phases) + " |")  # fmt: skip
    print("|" + "---|" * (5 + len(phases)))
    groups: dict[tuple[object, object], list[dict[str, object]]] = {}
    for report in reports:
        groups.setdefault((report["workspace_chunks"], report["scope"]), []).append(report)
    for (chunks, scope), items in groups.items():
        cells = [str(chunks), str(scope)] + [
            str(statistics.median(int(str(item[key])) for item in items))
            for key in ("candidates", "chunks", "context_tokens")
        ]
        for phase in phases:
            values = [item["ms"][phase] for item in items if phase in item["ms"]]  # type: ignore[index,operator]
            cells.append(
                f"{statistics.median(v['median'] for v in values):.1f} / "
                f"{max(v['p95'] for v in values):.1f}"
                if values
                else "-"
            )
        print("| " + " | ".join(cells) + " |")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--root", type=Path, default=user_cache_path("byod", appauthor=False) / "bench"
    )
    commands = parser.add_subparsers(dest="command", required=True)
    commands.add_parser("build").add_argument("--real-chunks", type=int, default=5000)
    matrix = commands.add_parser("run")
    matrix.add_argument("--runs", type=int, default=20)
    matrix.add_argument("--cases", default="0,3,7", help="Indexes of supported query cases")
    matrix.add_argument("--llm", metavar="OLLAMA_MODEL", help="Also measure time to first token")
    args = parser.parse_args()
    if args.command == "build":
        build(args.root, args.real_chunks)
    else:
        run(args.root, args.runs, [int(c) for c in args.cases.split(",")], args.llm)


if __name__ == "__main__":
    main()
