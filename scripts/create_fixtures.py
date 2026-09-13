"""Generate small, real format fixtures with authored, non-private study content."""

from pathlib import Path

import pymupdf
from docx import Document
from pptx import Presentation
from pptx.util import Inches

DESTINATION = Path(__file__).resolve().parents[1] / "tests" / "fixtures"
PARAGRAPH = (
    "Photosynthesis converts light energy into chemical energy in plant cells. "
    "Chlorophyll absorbs sunlight inside chloroplasts. Water and carbon dioxide "
    "are used to produce sugars and release oxygen. These reactions sustain "
    "plant growth and support food webs across terrestrial ecosystems."
)


def pdf_table(page: pymupdf.Page, top: float, start: int) -> None:
    rows = [["Process", "Location", "Result"]] + [
        [f"Reaction {i}", "Chloroplast membrane", "Energy transfer"]
        for i in range(start, start + 3)
    ]
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            rect = pymupdf.Rect(50 + j * 160, top + i * 30, 210 + j * 160, top + (i + 1) * 30)
            page.draw_rect(rect)
            page.insert_text((rect.x0 + 4, rect.y0 + 19), value, fontsize=10)


def main() -> None:
    DESTINATION.mkdir(parents=True, exist_ok=True)
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        page.insert_text((50, 60), "Plant Biology", fontsize=22)
        page.insert_text((50, 100), "Energy Conversion", fontsize=16)
        page.insert_textbox(pymupdf.Rect(50, 130, 540, 420), PARAGRAPH * 3, fontsize=11)
        pdf_table(page, 600, 1)
        page = pdf.new_page()
        pdf_table(page, 50, 4)
        page.insert_textbox(pymupdf.Rect(50, 220, 540, 700), PARAGRAPH * 4, fontsize=11)
        page = pdf.new_page()
        page.insert_text((50, 60), "Ecosystems", fontsize=22)
        page.insert_textbox(pymupdf.Rect(50, 100, 540, 700), PARAGRAPH * 4, fontsize=11)
        pdf.save(DESTINATION / "reading.pdf")

    with pymupdf.open() as scanned:
        image = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 40, 40), False)
        image.clear_with(200)
        for _ in range(2):
            page = scanned.new_page()
            page.insert_image(page.rect, pixmap=image)
        scanned.save(DESTINATION / "scanned.pdf")

    doc = Document()
    doc.add_heading("Plant Biology", 1)
    doc.add_paragraph(PARAGRAPH)
    doc.add_heading("Energy Conversion", 2)
    doc.add_paragraph(PARAGRAPH * 18)
    doc.add_heading("Chloroplasts", 3)
    doc.add_paragraph(PARAGRAPH)
    table = doc.add_table(rows=4, cols=2)
    for i, row in enumerate(table.rows):
        row.cells[0].text = f"Stage {i} of photosynthesis"
        row.cells[1].text = "Sunlight is absorbed and transformed into stored chemical energy."
    for _ in range(3):
        doc.add_paragraph(
            "Chlorophyll absorbs sunlight to power chemical reactions in plant cells.",
            style="List Bullet",
        )
    doc.save(str(DESTINATION / "notes.docx"))

    deck = Presentation()
    deck.core_properties.title = "Plant Biology"
    for number in range(1, 6):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        if number == 5:
            slide.shapes.add_shape(1, Inches(1), Inches(1), Inches(3), Inches(2))
            continue
        title = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(8), Inches(0.6))
        title.text_frame.text = f"Photosynthesis stage {number}"
        if number == 3:
            right = slide.shapes.add_textbox(Inches(5), Inches(1), Inches(4), Inches(3))
            right.text_frame.text = "Right column. " + PARAGRAPH
            left = slide.shapes.add_textbox(Inches(0.5), Inches(1), Inches(4), Inches(3))
            left.text_frame.text = "Left column. " + PARAGRAPH
        else:
            body = slide.shapes.add_textbox(Inches(0.5), Inches(1), Inches(8), Inches(3))
            body.text_frame.text = PARAGRAPH
        if number == 1:
            table = slide.shapes.add_table(3, 2, Inches(1), Inches(4), Inches(6), Inches(1)).table
            for row in table.rows:
                row.cells[0].text = "Chloroplast"
                row.cells[1].text = "Energy conversion"
        if number in {2, 4}:
            frame = slide.notes_slide.notes_text_frame
            if frame is not None:
                frame.text = (
                    "Explain the role of chlorophyll and connect light absorption "
                    "to sugar production."
                )
    deck.save(str(DESTINATION / "lecture.pptx"))


if __name__ == "__main__":
    main()
