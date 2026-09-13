from dataclasses import replace
from pathlib import Path

from pptx import Presentation

from byod.ingest.normalize import Block


def parse(path: Path) -> list[Block]:
    presentation = Presentation(str(path))
    deck_title = presentation.core_properties.title or path.stem
    blocks: list[Block] = []
    for number, slide in enumerate(presentation.slides, 1):
        if slide._element.get("show") == "0":
            continue
        locator = f"slide {number}"
        for shape in sorted(slide.shapes, key=lambda s: (s.top, s.left)):
            if shape.has_table:
                text = "\n".join(" | ".join(c.text for c in row.cells) for row in shape.table.rows)
                blocks.append(Block(text, locator, len(blocks), "pptx", "table", title=deck_title))
            elif shape.has_text_frame:
                is_title = shape == slide.shapes.title
                blocks.append(
                    Block(
                        shape.text_frame.text,
                        locator,
                        len(blocks),
                        "pptx",
                        "heading" if is_title else "body",
                        title=deck_title,
                    )
                )
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            if frame and frame.text.strip():
                blocks.append(
                    Block(frame.text, locator, len(blocks), "pptx", "notes", title=deck_title)
                )
    return [replace(block, unit_count=len(presentation.slides)) for block in blocks]
